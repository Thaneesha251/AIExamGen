import os
import hashlib
from typing import List, Optional
from pydantic import BaseModel
import docx
import pypdf


class ExtractionResult(BaseModel):
    text: str
    page_count: Optional[int] = None
    character_count: int
    warnings: List[str] = []
    checksum: str


class DocumentProcessor:
    """Document text extraction engine supporting PDF, DOCX, and TXT files."""

    ALLOWED_EXTENSIONS = {".pdf", ".docx", ".txt"}
    MAX_FILE_SIZE_BYTES = 25 * 1024 * 1024  # 25MB

    @staticmethod
    def calculate_checksum(file_bytes: bytes) -> str:
        return hashlib.sha256(file_bytes).hexdigest()

    def validate_file(self, filename: str, file_size: int):
        ext = os.path.splitext(filename)[1].lower()
        if ext not in self.ALLOWED_EXTENSIONS:
            raise ValueError(f"Unsupported file format '{ext}'. Allowed formats: .pdf, .docx, .txt")

        if file_size > self.MAX_FILE_SIZE_BYTES:
            raise ValueError(f"File size ({file_size / (1024*1024):.1f}MB) exceeds maximum limit of 25MB.")

    def extract_text(self, file_path: str, original_filename: str) -> ExtractionResult:
        if not os.path.exists(file_path):
            raise FileNotFoundError(f"File not found at path: {file_path}")

        file_size = os.path.getsize(file_path)
        self.validate_file(original_filename, file_size)

        with open(file_path, "rb") as f:
            file_bytes = f.read()
        checksum = self.calculate_checksum(file_bytes)

        ext = os.path.splitext(original_filename)[1].lower()

        if ext == ".txt":
            return self._extract_txt(file_bytes, checksum)
        elif ext == ".docx":
            return self._extract_docx(file_path, checksum)
        elif ext == ".pdf":
            return self._extract_pdf(file_path, checksum)
        else:
            raise ValueError(f"Unsupported extension: {ext}")

    def _extract_txt(self, file_bytes: bytes, checksum: str) -> ExtractionResult:
        warnings = []
        try:
            text = file_bytes.decode("utf-8")
        except UnicodeDecodeError:
            text = file_bytes.decode("utf-8", errors="replace")
            warnings.append("Encoding issue detected: Non-UTF8 characters replaced.")

        return ExtractionResult(
            text=text.strip(),
            page_count=1,
            character_count=len(text.strip()),
            warnings=warnings,
            checksum=checksum
        )

    def _extract_docx(self, file_path: str, checksum: str) -> ExtractionResult:
        warnings = []
        try:
            doc = docx.Document(file_path)
            paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
            
            # Extract text inside tables as well
            for table in doc.tables:
                for row in table.rows:
                    row_text = " | ".join([cell.text.strip() for cell in row.cells if cell.text.strip()])
                    if row_text:
                        paragraphs.append(row_text)

            text = "\n\n".join(paragraphs)
            if not text.strip():
                warnings.append("Document appears empty or contains no extractable text paragraphs.")

            return ExtractionResult(
                text=text.strip(),
                page_count=None,
                character_count=len(text.strip()),
                warnings=warnings,
                checksum=checksum
            )
        except Exception as e:
            raise RuntimeError(f"Failed to extract text from DOCX file: {str(e)}")

    def _extract_pdf(self, file_path: str, checksum: str) -> ExtractionResult:
        warnings = []
        page_texts = []
        page_count = 0

        try:
            reader = pypdf.PdfReader(file_path)
            page_count = len(reader.pages)

            for idx, page in enumerate(reader.pages, start=1):
                extracted = page.extract_text() or ""
                if extracted.strip():
                    page_texts.append(f"--- PAGE {idx} ---\n{extracted.strip()}")

            full_text = "\n\n".join(page_texts).strip()

            # Scanned image PDF detection
            if page_count > 0 and len(full_text) < 20:
                warnings.append(
                    "TEXT_EXTRACTION_UNAVAILABLE: PDF appears to be a scanned image or contains vector text. OCR processing will be required."
                )

            return ExtractionResult(
                text=full_text,
                page_count=page_count,
                character_count=len(full_text),
                warnings=warnings,
                checksum=checksum
            )
        except Exception as e:
            raise RuntimeError(f"Failed to extract text from PDF file: {str(e)}")
